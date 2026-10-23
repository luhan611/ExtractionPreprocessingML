"""Integration acceptance checks against the local papers (not an accuracy benchmark)."""
import json
from pathlib import Path

import pdfplumber
from docx import Document
from openpyxl import load_workbook

from lnp_preprocess.tools import DocumentTools


def main():
    report = []
    for paper in sorted(Path('Paper').iterdir()):
        bundle_path = Path('processed') / paper.name / 'document_bundle.json'
        if not bundle_path.exists():
            raise SystemExit(f'Missing bundle: {bundle_path}; parse it first')
        tools = DocumentTools(bundle_path)
        bundle = tools.bundle
        assert all(d['status'] == 'parsed' for d in bundle['documents'])
        for key in ('documents', 'chunks', 'tables', 'figures'):
            ids = [x['id'] for x in bundle[key]]
            assert len(ids) == len(set(ids)), (paper, key, 'duplicate IDs')
        summary = {'paper_id': paper.name, 'pdf_pages_preserved': 0,
                   'word_tables': [], 'excel_sheets': []}
        for doc in bundle['documents']:
            tables = [t for t in bundle['tables'] if t['document_id'] == doc['id']]
            if doc['format'] == 'xlsx':
                book = load_workbook(paper / doc['filename'], data_only=False)
                try:
                    for sheet, table in zip(book, tables, strict=True):
                        expected = {c.coordinate: c.value for row in sheet for c in row if c.value is not None}
                        actual = {c['coordinate']: c['raw_value'] for c in table['cells']}
                        assert actual == expected, (paper, doc['filename'], sheet.title)
                        summary['excel_sheets'].append({'file': doc['filename'], 'sheet': sheet.title,
                                                       'nonempty_rows': len({c['row'] for c in table['cells']}),
                                                       'nonempty_cells': len(actual)})
                finally:
                    book.close()
            elif doc['format'] == 'docx':
                word = Document(paper / doc['filename'])
                for source, table in zip(word.tables, tables, strict=True):
                    expected = [c.text for row in source.rows for c in row.cells]
                    assert [c['raw_value'] for c in table['cells']] == expected
                    summary['word_tables'].append({'file': doc['filename'], 'rows': table['row_count']})
            elif doc['format'] == 'pdf':
                with pdfplumber.open(paper / doc['filename']) as pdf:
                    pages = [f for f in bundle['figures'] if f['document_id'] == doc['id'] and f['kind'] == 'page_image']
                    assert {f['locator']['page'] for f in pages} == set(range(1, len(pdf.pages) + 1))
                    summary['pdf_pages_preserved'] += len(pages)
        for figure in bundle['figures']:
            assert (bundle_path.parent / figure['asset_path']).exists()
        # All parsed text is recoverable via budgeted pagination.
        offset, ids = 0, []
        while True:
            page = tools.get_full_text(offset=offset, max_chars=5000)
            ids.extend(c['id'] for c in page['items'])
            if not page['truncated']:
                break
            offset = page['next_offset']
        assert ids == [c['id'] for c in bundle['chunks']]
        summary['issues'] = [i['code'] for i in bundle['issues']]
        report.append(summary)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
