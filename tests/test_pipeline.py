import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image

from lnp_preprocess.parser import parse_paper
from lnp_preprocess.tools import DocumentTools


def write_pdf(path):
    # Tiny valid text PDF fixture, no production parser implementation here.
    stream = b'BT /F1 12 Tf 20 100 Td (Figure 1. ALC-0314 sample.) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 250 150] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    data, offsets = b'%PDF-1.4\n', [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
    start = len(data)
    data += b'xref\n0 6\n0000000000 65535 f \n'
    data += b''.join(f'{n:010d} 00000 n \n'.encode() for n in offsets[1:])
    data += f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()
    path.write_bytes(data)


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.paper = self.root / 'P_TEST'
        self.paper.mkdir()
        image = self.root / 'sample.png'
        Image.new('RGB', (20, 20), 'white').save(image)
        word = Document()
        word.add_paragraph('Before Table S1. ' + 'long paragraph ' * 30)
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged header'
        table.cell(1, 0).text = 'ALC-0314'
        table.cell(1, 1).text = '12.5'
        word.add_paragraph('After table')
        word.add_picture(str(image))
        word.save(self.paper / 'si.docx')
        book = Workbook()
        sheet = book.active
        sheet.title = 'Results'
        sheet.append(['Name', 'Size', 'Computed'])
        sheet.append(['CEE-E12', 12.5, '=B2*2'])
        sheet.merge_cells('A3:B3')
        sheet['A3'] = 'merged'
        sheet['V55'].number_format = '0.00'  # Formatting must not create data.
        book.save(self.paper / 'si.xlsx')
        write_pdf(self.paper / 'main.pdf')
        self.output = self.root / 'processed'
        self.bundle = parse_paper(self.paper, self.output, max_chars=150)
        self.tools = DocumentTools(self.output / 'document_bundle.json')

    def tearDown(self):
        self.temp.cleanup()

    def test_documents_and_anchors(self):
        self.assertTrue(all(d['status'] == 'parsed' for d in self.bundle['documents']))
        ids = [x['id'] for key in ('chunks', 'tables', 'figures') for x in self.bundle[key]]
        self.assertEqual(len(ids), len(set(ids)))
        for key in ('chunks', 'tables', 'figures'):
            for item in self.bundle[key]:
                self.assertIn(item['document_id'], self.tools.documents)
        for figure in self.bundle['figures']:
            self.assertTrue((self.output / figure['asset_path']).is_file())
        pdf_chunk = next(x for x in self.bundle['chunks'] if x['kind'] == 'pdf_layout_block')
        region = self.tools.get_source_region(pdf_chunk['id'])
        self.assertEqual(region['locator']['page'], 1)
        self.assertEqual(region['assets'][0]['kind'], 'page_image')

    def test_word_order_and_merged_cells(self):
        word = next(d for d in self.bundle['documents'] if d['format'] == 'docx')
        chunks = self.tools._document_chunks(word['id'])
        positions = [c['locator']['body_index'] for c in chunks]
        self.assertEqual(positions, sorted(positions))
        table = next(t for t in self.bundle['tables'] if t['document_id'] == word['id'])
        self.assertEqual(table['cells'][1]['merged_origin'], [1, 1])
        self.assertEqual(chunks[-1]['text'], 'After table')
        self.assertTrue(any(f['kind'] == 'embedded_image' for f in self.bundle['figures']))

    def test_excel_values_formulas_and_empty_tail(self):
        table = next(t for t in self.bundle['tables'] if 'sheet' in t['locator'])
        cells = {c['coordinate']: c for c in table['cells']}
        self.assertEqual(cells['A2']['raw_value'], 'CEE-E12')
        self.assertEqual(cells['B2']['raw_value'], 12.5)
        self.assertEqual(cells['C2']['formula'], '=B2*2')
        self.assertIsNone(cells['C2']['cached_value'])
        self.assertEqual(table['row_count'], 3)
        self.assertNotIn('V55', cells)
        self.assertEqual(table['merged_ranges'], ['A3:B3'])
        self.assertTrue(any(x['code'] == 'formula_cache_missing' for x in self.bundle['issues']))
        page = self.tools.get_table(table['id'], limit=2)['cell_page']
        self.assertTrue(page['truncated'])
        self.assertEqual(page['next_offset'], 2)

    def test_pagination_no_text_lost(self):
        offset, collected = 0, []
        while True:
            result = self.tools.get_full_text(offset=offset, max_chars=200)
            collected.extend(result['items'])
            if not result['truncated']:
                break
            offset = result['next_offset']
        self.assertEqual(collected, self.bundle['chunks'])
        with self.assertRaises(ValueError):
            self.tools.get_full_text(max_chars=1)
        with self.assertRaises(KeyError):
            self.tools.get_full_text(document_id='does-not-exist')
        matches = self.tools.search_document('ALC-0314', limit=1)
        self.assertTrue(matches['truncated'])
        anchor = matches['items'][0]['id']
        self.assertTrue(self.tools.get_chunks([anchor], neighbors=1)['items'])
        self.assertEqual(self.tools.resolve_reference('Figure 1')['status'], 'unresolved')

    def test_reproducible_ids_and_failed_input(self):
        second = parse_paper(self.paper, self.root / 'second', max_chars=150)
        self.assertEqual(self.bundle, second)
        (self.paper / 'broken.docx').write_bytes(b'not a zip')
        third = parse_paper(self.paper, self.root / 'third', max_chars=150)
        self.assertEqual(next(d for d in third['documents'] if d['filename'] == 'broken.docx')['status'], 'failed')
        self.assertTrue(any(i['code'] == 'document_parse_failed' for i in third['issues']))


if __name__ == '__main__':
    unittest.main()
